from decimal import Decimal

from django.db.models import F, Sum, DecimalField, ExpressionWrapper, Q
from django.shortcuts import render, redirect, get_object_or_404

from .models import Category, Product
from .forms import ProductForm


def dashboard(request):
    active_products = Product.objects.filter(is_active=True)

    stock_value_expression = ExpressionWrapper(
        F("stock_quantity") * F("cost_price"),
        output_field=DecimalField(max_digits=20, decimal_places=2),
    )

    context = {
        "total_products": active_products.count(),
        "out_of_stock_count": active_products.filter(
            stock_quantity=0
        ).count(),
        "total_categories": Category.objects.count(),
        "total_stock_units": active_products.aggregate(
            total=Sum("stock_quantity")
        )["total"] or 0,
        "inventory_cost_value": active_products.aggregate(
            total=Sum(stock_value_expression)
        )["total"] or Decimal("0.00"),
        "low_stock_products": active_products.filter(
            stock_quantity__lte=F("reorder_level")
        ).order_by("stock_quantity", "name")[:10],
        "recent_products": active_products.order_by("-created_at")[:10],
    }
    return render(request, "products/dashboard.html", context)


def product_list(request):
    query = request.GET.get("q", "").strip()
    products = Product.objects.select_related("category").order_by("name")

    if query:
        products = products.filter(
            Q(name__icontains=query) | Q(sku__icontains=query)
        )

    return render(request, "products/product_list.html", {
        "products": products,
        "query": query,
    })


def product_create(request):
    form = ProductForm(request.POST or None)

    if request.method == "POST" and form.is_valid():
        form.save()
        return redirect("products:product_list")

    return render(request, "products/product_form.html", {
        "form": form,
        "heading": "Add Product",
    })


def product_edit(request, pk):
    product = get_object_or_404(Product, pk=pk)
    form = ProductForm(request.POST or None, instance=product)

    if request.method == "POST" and form.is_valid():
        form.save()
        return redirect("products:product_list")

    return render(request, "products/product_form.html", {
        "form": form,
        "heading": "Edit Product",
    })


def product_delete(request, pk):
    product = get_object_or_404(Product, pk=pk)

    if request.method == "POST":
        product.delete()
        return redirect("products:product_list")

    return render(request, "products/product_confirm_delete.html", {
        "product": product,
    })


def category_list(request):
    query = request.GET.get("q", "").strip()
    categories = Category.objects.order_by("name")

    if query:
        categories = categories.filter(name__icontains=query)

    return render(request, "products/category_list.html", {
        "categories": categories,
        "query": query,
    })


def category_create(request):
    from django import forms

    class CategoryForm(forms.ModelForm):
        class Meta:
            model = Category
            fields = ["name"]

    form = CategoryForm(request.POST or None)

    if request.method == "POST" and form.is_valid():
        form.save()
        return redirect("products:category_list")

    return render(request, "products/category_form.html", {
        "form": form,
        "heading": "Add Category",
    })


def category_edit(request, pk):
    from django import forms

    category = get_object_or_404(Category, pk=pk)

    class CategoryForm(forms.ModelForm):
        class Meta:
            model = Category
            fields = ["name"]

    form = CategoryForm(request.POST or None, instance=category)

    if request.method == "POST" and form.is_valid():
        form.save()
        return redirect("products:category_list")

    return render(request, "products/category_form.html", {
        "form": form,
        "heading": "Edit Category",
    })


def category_delete(request, pk):
    category = get_object_or_404(Category, pk=pk)

    if request.method == "POST":
        try:
            category.delete()
            return redirect("products:category_list")
        except Exception:
            from django.contrib import messages
            messages.error(
                request,
                "This category cannot be deleted because products may use it."
            )

    return render(request, "products/category_confirm_delete.html", {
        "category": category,
    })


from django.contrib import messages
from django.db import transaction
from django.core.exceptions import ValidationError
from .models import StockTransaction
from .forms import StockTransactionForm


def stock_transaction_create(request):
    form = StockTransactionForm(request.POST or None)

    if request.method == "POST" and form.is_valid():
        try:
            with transaction.atomic():
                product_id = form.cleaned_data["product"].pk
                product = Product.objects.select_for_update().get(
                    pk=product_id
                )
                movement = form.save(commit=False)
                movement.product = product
                quantity = movement.quantity

                if movement.transaction_type == "OUT":
                    if quantity > product.stock_quantity:
                        form.add_error(
                            "quantity",
                            "Not enough stock available.",
                        )
                        raise ValidationError("Insufficient stock.")

                    product.stock_quantity -= quantity

                elif movement.transaction_type == "IN":
                    product.stock_quantity += quantity

                elif movement.transaction_type == "ADJUSTMENT":
                    product.stock_quantity = quantity

                product.save(update_fields=["stock_quantity"])
                movement.save()

            messages.success(request, "Stock updated successfully.")
            return redirect("products:stock_transaction_create")

        except ValidationError:
            pass

    transactions = StockTransaction.objects.select_related(
        "product"
    ).all()[:20]

    return render(request, "products/stock_transactions.html", {
        "form": form,
        "transactions": transactions,
    })


from django.db import transaction
from django.contrib import messages
from django.core.exceptions import ValidationError
from .forms import SaleForm
from .models import Sale


def sale_create(request):
    form = SaleForm(request.POST or None)

    if request.method == "POST" and form.is_valid():
        try:
            with transaction.atomic():
                product_id = form.cleaned_data["product"].pk
                product = Product.objects.select_for_update().get(
                    pk=product_id
                )
                quantity = form.cleaned_data["quantity"]

                if quantity > product.stock_quantity:
                    form.add_error(
                        "quantity",
                        "Not enough stock available.",
                    )
                    raise ValidationError("Insufficient stock.")

                sale = form.save(commit=False)
                sale.product = product
                sale.unit_price = form.cleaned_data["unit_price"]
                product.stock_quantity -= quantity
                product.save(update_fields=["stock_quantity"])
                sale.save()

            messages.success(request, "Sale recorded successfully.")
            return redirect("products:sale_create")

        except ValidationError:
            pass

    sales = Sale.objects.select_related("product")[:20]

    return render(request, "products/sales.html", {
        "form": form,
        "sales": sales,
    })


def sales_report(request):
    from decimal import Decimal
    from django.db.models import Sum, F, ExpressionWrapper, DecimalField
    from django.utils import timezone
    from .models import Sale

    money_field = DecimalField(max_digits=24, decimal_places=2)

    revenue_expression = ExpressionWrapper(
        F("quantity") * F("unit_price"),
        output_field=money_field,
    )

    profit_expression = ExpressionWrapper(
        F("quantity") * (F("unit_price") - F("product__cost_price")),
        output_field=money_field,
    )

    all_sales = Sale.objects.select_related("product")
    today = timezone.localdate()
    today_sales = all_sales.filter(created_at__date=today)

    totals = all_sales.aggregate(
        revenue=Sum(revenue_expression),
        profit=Sum(profit_expression),
        units=Sum("quantity"),
    )

    today_revenue = today_sales.aggregate(
        revenue=Sum(revenue_expression)
    )["revenue"] or Decimal("0.00")

    context = {
        "total_sales": all_sales.count(),
        "total_units": totals["units"] or 0,
        "total_revenue": totals["revenue"] or Decimal("0.00"),
        "estimated_profit": totals["profit"] or Decimal("0.00"),
        "today_revenue": today_revenue,
        "sales": all_sales[:100],
    }

    return render(request, "products/sales_report.html", context)


def sales_export_excel(request):
    from django.http import HttpResponse
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from .models import Sale

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sales Report"

    sheet.append([
        "Date", "SKU", "Product", "Quantity",
        "Unit Price", "Total Revenue", "Estimated Profit",
    ])

    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(
            fill_type="solid", fgColor="2563EB"
        )

    sales = Sale.objects.select_related("product").order_by("-created_at")

    for sale in sales:
        sheet.append([
            sale.created_at.strftime("%Y-%m-%d %H:%M"),
            sale.product.sku,
            sale.product.name,
            sale.quantity,
            float(sale.unit_price),
            float(sale.total_amount),
            float(sale.estimated_profit),
        ])

    for column in sheet.columns:
        letter = column[0].column_letter
        sheet.column_dimensions[letter].width = min(
            max(max(len(str(c.value or "")) for c in column) + 2, 12),
            30,
        )

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = (
        'attachment; filename="sales_report.xlsx"'
    )
    workbook.save(response)
    return response


def sales_export_pdf(request):
    from django.http import HttpResponse
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import landscape, A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import (
        SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
    )
    from .models import Sale

    response = HttpResponse(content_type="application/pdf")
    response["Content-Disposition"] = (
        'attachment; filename="sales_report.pdf"'
    )

    document = SimpleDocTemplate(
        response,
        pagesize=landscape(A4),
        rightMargin=25,
        leftMargin=25,
        topMargin=25,
        bottomMargin=25,
    )

    styles = getSampleStyleSheet()
    elements = [
        Paragraph("Product ERP - Sales Report", styles["Title"]),
        Spacer(1, 15),
    ]

    sales = Sale.objects.select_related(
        "product"
    ).order_by("-created_at")

    rows = [[
        "Date", "SKU", "Product", "Qty",
        "Unit Price", "Revenue", "Est. Profit",
    ]]

    total_revenue = 0
    total_profit = 0

    for sale in sales:
        revenue = sale.total_amount
        profit = sale.estimated_profit
        total_revenue += revenue
        total_profit += profit

        rows.append([
            sale.created_at.strftime("%Y-%m-%d"),
            str(sale.product.sku),
            Paragraph(str(sale.product.name), styles["BodyText"]),
            str(sale.quantity),
            f"{sale.unit_price:,.2f}",
            f"{revenue:,.2f}",
            f"{profit:,.2f}",
        ])

    rows.append([
        "", "", "TOTAL", "", "",
        f"{total_revenue:,.2f}",
        f"{total_profit:,.2f}",
    ])

    table = Table(rows, repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2563EB")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.lightgrey),
        ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#DBEAFE")),
        ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
    ]))

    elements.append(table)
    elements.append(Spacer(1, 12))
    elements.append(Paragraph(
        "Estimated profit uses current product cost prices, "
        "not necessarily the cost at the time of sale.",
        styles["Italic"],
    ))

    document.build(elements)
    return response


def sales_dashboard(request):
    from django.db.models import Sum, Count
    from django.shortcuts import render
    from .models import Sale, Product

    sales = Sale.objects.select_related("product").order_by("-created_at")
    total_sales = sales.count()
    total_revenue = sum((sale.total_amount for sale in sales), 0)
    total_profit = sum((sale.estimated_profit for sale in sales), 0)

    context = {
        "sales": sales[:10],
        "total_sales": total_sales,
        "total_revenue": total_revenue,
        "total_profit": total_profit,
        "total_products": Product.objects.count(),
        "total_units_sold": sales.aggregate(
            total=Sum("quantity")
        )["total"] or 0,
    }

    return render(request, "products/sales_dashboard.html", context)
